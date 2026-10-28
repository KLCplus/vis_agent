"""Create a Chinese classroom report using generated chapters and figures."""
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from generate_docs import main, ROOT


def create():
    chapters=main();import json;data=json.loads((ROOT/'data/processed/dashboard.json').read_text());doc=Document();section=doc.sections[0]
    section.page_width=Cm(21);section.page_height=Cm(29.7)
    section.top_margin=Cm(1.8);section.bottom_margin=Cm(1.8);section.left_margin=Cm(2);section.right_margin=Cm(2)
    for style_name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Caption']:
        st=doc.styles[style_name];st.font.name='Arial Unicode MS';st.font.color.rgb=RGBColor(0,0,0)
        st.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Arial Unicode MS')
    normal=doc.styles['Normal'];normal.font.size=Pt(10.5);normal.paragraph_format.line_spacing=1.12;normal.paragraph_format.space_after=Pt(5);normal.paragraph_format.keep_with_next=False
    doc.styles['Title'].font.size=Pt(25);doc.styles['Heading 1'].font.size=Pt(18);doc.styles['Heading 2'].font.size=Pt(11.5)
    doc.styles['Caption'].font.size=Pt(9)
    for st in doc.styles:
        for border in list(st.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
    for st in ['Title','Subtitle','Heading 1','Heading 2']:
        doc.styles[st].font.color.rgb=RGBColor(0,0,0)
    doc.styles['Heading 2'].paragraph_format.space_before=Pt(8)
    doc.styles['Heading 2'].paragraph_format.space_after=Pt(3)
    doc.styles['Caption'].paragraph_format.keep_with_next=False
    figure_number=0
    doc.add_paragraph('多 Agent 异常行为与信息传播可视分析',style='Title')
    doc.add_paragraph('VAST 2026 MC1 课程第一次作业',style='Subtitle')
    doc.add_paragraph('本报告完成数据下载、数据理解、数据过滤和数据处理四个设计步骤，并实现可追溯的交互分析原型。全部统计取自官方 MC1 数据，清洗结果、图表和证据可以通过代码重新生成。')
    doc.add_paragraph(f"数据包含 {data['profile']['agent_count']} 个 Agent、{data['profile']['message_count']} 条消息与 {data['profile']['round_count']} 个非等间隔轮次。主要质量风险来自无效及指向未来的回复引用；第一阶段保留原始记录并阻止这些引用构成错误传播线索。正式解禁前的公开记录列为候选，授权情况需继续核验。")
    doc.add_paragraph('小组名称 ［待填写］\n组员姓名与学号 ［待填写］\n课程与教师 ［待填写］\n提交日期 ［待填写］')
    doc.add_heading('阅读与复现',level=2)
    doc.add_paragraph('本报告正文为课程初稿。代码入口、数据获取和本地启动命令见仓库 README。原始数据不入库，标准表和页面快照由 scripts/preprocess.py 重新生成。事实记录、回复推断及尚未确定的传播关系分别标注。')
    doc.add_paragraph('官方任务 https://vast-challenge.github.io/2026/MC1.html\n官方数据 https://github.com/vast-challenge/2026-data\n项目仓库 https://github.com/KLCplus/vis_agent',style='Caption')
    for idx,(filename,title,sections,figs) in enumerate(chapters,1):
        doc.add_page_break();doc.add_heading(f'{idx} {title}',level=1)
        for heading,body in sections:
            doc.add_heading(title+'的'+heading,level=2)
            # Keep prose readable; convert Markdown citation syntax to human-readable text.
            import re
            body=re.sub(r'\[([^]]+)\]\(([^)]+)\)',lambda m:m[1]+' '+m[2],body)
            doc.add_paragraph(body)
        if idx in [2,5]:
            doc.add_page_break();doc.add_heading(title+'统计图表',level=2)
        # Dashboard screenshot is tall. Use top viewport excerpt generated from actual screenshot.
        for f in figs:
            path=ROOT/'reports/figures'/(f+'.png')
            if f=='08_dashboard' and (ROOT/'reports/figures/08_dashboard_overview.png').exists():path=ROOT/'reports/figures/08_dashboard_overview.png'
            if path.exists():
                paragraph=doc.add_paragraph();paragraph.paragraph_format.keep_with_next=True
                paragraph.add_run().add_picture(str(path),width=Inches(4.7 if f=='08_dashboard' else 5.7 if f=='06_network' else 6.3))
                figure_number+=1
                doc.add_paragraph(f'图 {figure_number} {title}  来源为程序生成的真实数据图表',style='Caption')
    for paragraph in doc.paragraphs:
        if paragraph.style.name not in ['Heading 1','Heading 2']:
            paragraph.paragraph_format.keep_with_next=False
        for run in paragraph.runs:
            run.font.name='Arial Unicode MS'
            rf=run._element.get_or_add_rPr().rFonts
            for attr in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:
                rf.attrib.pop(qn('w:'+attr),None)
            rf.set(qn('w:eastAsia'),'Arial Unicode MS')
    out=ROOT/'reports/第一次作业_初稿.docx';doc.save(out);print(out)

if __name__=='__main__':create()
