"""Create the course Word report with retained intermediate evidence and real figures."""
import csv
import json
import platform
import re
from pathlib import Path
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
EVIDENCE = ROOT / 'reports/evidence'
OUTPUT = ROOT / 'reports/第一次作业_数据处理与中间证据.docx'


def create():
    data = json.loads((ROOT/'data/processed/dashboard.json').read_text())
    evidence = json.loads((EVIDENCE/'summary.json').read_text())
    examples = json.loads((EVIDENCE/'06_before_after_examples.json').read_text())
    markers = json.loads((EVIDENCE/'11_key_evidence_records.json').read_text())
    verification = json.loads((EVIDENCE/'12_verification_runs.json').read_text())
    passed_tests = int(next(re.search(r'(\d+) passed', run['stdout']).group(1) for run in verification if re.search(r'(\d+) passed', run['stdout'])))
    def csv_rows(name):
        with (EVIDENCE/name).open(encoding='utf-8-sig') as stream: return list(csv.DictReader(stream))
    doc = Document(); sec=doc.sections[0]
    sec.page_width=Cm(21);sec.page_height=Cm(29.7)
    sec.top_margin=Cm(1.7);sec.bottom_margin=Cm(1.7);sec.left_margin=Cm(2);sec.right_margin=Cm(2)
    for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','Caption']:
        style=doc.styles[name];style.font.name='Arial Unicode MS';style.font.color.rgb=RGBColor(0,0,0)
        style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Arial Unicode MS')
        style.paragraph_format.keep_with_next=False
    normal=doc.styles['Normal'];normal.font.size=Pt(10.5);normal.paragraph_format.line_spacing=1.2;normal.paragraph_format.space_after=Pt(6)
    normal.paragraph_format.widow_control=True
    for name,size in [('Title',22),('Subtitle',12),('Heading 1',16),('Heading 2',12),('Heading 3',11),('Caption',8.5)]:doc.styles[name].font.size=Pt(size)
    for name in ['Heading 1','Heading 2','Heading 3']:
        style=doc.styles[name];style.font.bold=True;style.paragraph_format.keep_with_next=True;style.paragraph_format.space_before=Pt(9);style.paragraph_format.space_after=Pt(6)
    for style in doc.styles:
        for border in list(style.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
    for name in ['Normal','Caption','Subtitle']:
        doc.styles[name].font.bold=False;doc.styles[name].font.italic=False
    # Page numbers aid course review; no ornamental header.
    footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    figure_count=0;table_count=0;source_index=[];outline=[]
    def paragraph(text,style=None):
        p=doc.add_paragraph(str(text),style=style);return p
    def page(title):
        if doc.paragraphs and any(p.text.strip() for p in doc.paragraphs):doc.add_page_break()
        doc.add_heading(title,level=1);outline.append(title)
    def heading(text):doc.add_heading(text,level=2)
    def code(text):
        for line in text.splitlines():
            p=doc.add_paragraph();p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.0
            run=p.add_run(line);run.font.name='Courier New';run.font.size=Pt(8.5)
        paragraph('')
    def figure(name,caption,width=6.35):
        nonlocal figure_count
        path=ROOT/'reports/figures'/name
        assert path.exists(),path
        p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True
        p.add_run().add_picture(str(path),width=Inches(width))
        figure_count+=1
        p=paragraph(f'图 {figure_count}  {caption}','Caption');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        source_index.append({'number':figure_count,'figure':str(path.relative_to(ROOT)),'caption':caption})
    def table(headers, rows, widths, caption):
        nonlocal table_count
        table_count+=1;paragraph(f'表 {table_count}  {caption}','Caption')
        t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
        for i,width in enumerate(widths):t.columns[i].width=Cm(width)
        for cell,header in zip(t.rows[0].cells,headers):cell.text=str(header)
        repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
        for row in rows:
            cells=t.add_row().cells
            for cell,value in zip(cells,row):cell.text=str(value)
        for ri,row in enumerate(t.rows):
            for ci,cell in enumerate(row.cells):
                cell.width=Cm(widths[ci]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                pr=cell._tc.get_or_add_tcPr();borders=OxmlElement('w:tcBorders')
                for edge in ['top','left','bottom','right']:
                    e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
                pr.append(borders)
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E6EAF1' if ri==0 else 'F7F8FA' if ri%2==0 else 'FFFFFF');pr.append(shade)
                margins=OxmlElement('w:tcMar')
                for edge in ['top','bottom','left','right']:
                    margin=OxmlElement('w:'+edge);margin.set(qn('w:w'),'75');margin.set(qn('w:type'),'dxa');margins.append(margin)
                pr.append(margins)
                for p in cell.paragraphs:
                    p.paragraph_format.space_after=Pt(2);p.paragraph_format.space_before=Pt(2);p.paragraph_format.line_spacing=1.05
                    for run in p.runs:run.font.size=Pt(9);run.bold=ri==0
        paragraph('')
    total=data['profile']['message_count'];public=sum(message['scope']=='external' for message in data['messages'])
    q=evidence['quality']

    doc.add_paragraph('多 Agent 异常行为与信息传播可视分析',style='Title')
    doc.add_paragraph('第一次作业 数据处理实验报告',style='Subtitle')
    paragraph('课程名称  ［待填写］\n小组名称  ［待填写］\n姓名及学号  ［待填写］\n指导教师  ［待填写］\n完成日期  ［待填写］')
    heading('本次实验内容')
    paragraph('本次实验以 VAST Challenge 2026 Mini-Challenge 1 的官方数据为对象，主要完成数据下载、数据理解、数据过滤和数据处理四个步骤。在此基础上，再把处理结果绘制成图表，查看不同 Agent 的沟通方式及与 HarborCrest 并购相关的记录。')
    paragraph('报告按实际操作顺序整理。每一步都保留了输入、处理方法、输出结果和对应的中间文件，便于之后写课程作业或重新检查数据时继续使用。')
    table(['检查项目','本次实际结果'],[['消息与角色',f'{total} 条消息，{len(data["agents"])} 个 Agent'],['轮次与时间',f'{len(data["rounds"])} 个轮次，2046年5月17日至6月5日'],['处理结果','8 张标准化表，原消息全部保留'],['分析结果',f'{len(data["candidates"])} 条异常候选，需结合原文复核']],[4.2,12.7],'实验数据概况')
    paragraph('官方数据来源 https://github.com/vast-challenge/2026-data\n官方任务说明 https://vast-challenge.github.io/2026/MC1.html\n项目仓库 https://github.com/KLCplus/vis_agent',style='Caption')

    page('一 实验目的与实验方法')
    paragraph('这次实验首先需要弄清楚原始文件到底记录了什么。如果直接根据比赛背景去画图，就容易把角色名称、消息关系和时间间隔理解错。因此，我先检查数据结构，再进行字段整理，最后才计算统计结果。')
    paragraph('本次希望解决的问题有三个：数据能否被完整读取，处理后的每条记录能否回到原始消息，以及图表中的数量能否由处理结果重新计算。敏感主题和异常行为先用规则初筛，不使用大模型 API。')
    figure('01_pipeline.png','从官方文件到标准化数据和图表的处理步骤')
    table(['项目','采用的方法'],[['数据读取','Python JSON 与 CSV'],['统计处理','Pandas 和标准 Python 统计'],['通信关系','接收人映射与回复引用检查，NetworkX 绘图'],['交互图表','React、TypeScript、Apache ECharts'],['数据服务','FastAPI，只读结构化文件']],[3.5,13.4],'实验工具与方法')
    paragraph('图中的每一步都有文件输出。这样即使后面的图表需要修改，也不必重新猜测前面是如何处理的。完整流程可以通过脚本重复执行。')

    page('二 数据下载与原始文件检查')
    paragraph('首先从官方仓库下载 VAST_Challenge_2026_MC1.zip，再解压检查文件。压缩包并不是只有一个 JSON，其中还带有数据说明、作答表和 macOS 生成的辅助文件。这里先把所有条目列出来，再区分哪些文件参与分析。')
    inventory=csv_rows('01_archive_inventory.csv')
    table(['文件或类别','数量','大小'],[['MC1_final_00.json',1,f'{next(int(r["bytes"]) for r in inventory if r["name"].endswith(".json")):,} B'],['mc1 data description.md',1,f'{next(int(r["bytes"]) for r in inventory if r["name"].endswith(".md")):,} B'],['Answer Sheet.htm',1,f'{next(int(r["bytes"]) for r in inventory if r["name"].endswith(".htm")):,} B'],['macOS 辅助文件',evidence['archive_metadata_files'],f'{sum(int(r["bytes"]) for r in inventory if r["category"]=="系统元数据"):,} B']],[9.8,2.0,5.1],'压缩包文件清单汇总')
    figure('evidence/11_archive.png','解压条目的大小对比')
    paragraph(f'可以看到，真正承载实验数据的是 JSON 文件。压缩包共有 {evidence["archive_file_count"]} 个文件条目，其中 {evidence["archive_semantic_files"]} 个是需要阅读的语义文件，另外 {evidence["archive_metadata_files"]} 个属于系统元数据，不作为消息记录处理。')
    paragraph('中间证据：reports/evidence/01_archive_inventory.csv。原始压缩包保存在 data/raw，下载脚本支持重新获取，也支持从已有官方压缩包离线解压。',style='Caption')

    page('三 理解实际 JSON 结构')
    paragraph('解压后，我先查看 JSON 的顶层结构。实际文件只有 rounds 这个顶层字段，每个轮次再包含环境信息、消息列表和参与者列表。这里与数据说明中的示例并不完全一致，消息是直接放在 communications 中的，因此后面的程序以真实文件为准。')
    figure('evidence/12_schema.png','原始文件中的主要层次和实际记录数量')
    table(['字段','读取后的含义'],[['hour','轮次起点，不是每条消息自己的发送时间'],['environment_context','市场快照、外部事件、人员不可用等环境'],['communications','逐条消息、正文、内部状态与消息引用'],['participants','该轮参与角色和声明的动作'],['internal_state','reacting、rationalizing、deliberating 三种状态流']],[5.2,11.7],'本次解析使用的主要字段')
    paragraph(f'这一步共递归检查了 {data["profile"]["field_paths"]} 个字段路径。字段字典保留每种类型出现的次数及 null 数量，后面发现异常值时可以据此定位，而不是只看某一条示例记录。')
    paragraph('中间证据：data/reports/data_dictionary.csv；reports/evidence/02_message_field_audit.csv。',style='Caption')

    page('四 Agent 角色与消息规模')
    paragraph('接下来按照原始 agent_id 统计每个 Agent 的消息量。这里不根据正文语气判断“实际是谁在说话”，因为这样会改变原始记录的归属。角色标签只用于显示，原 ID 始终保留。')
    figure('02_agent_messages.png','各 Agent 发出的消息数量')
    table(['实际 Agent ID','原始职能 role','消息数'],[[a['agent_id'],a['role'],a['message_count']] for a in data['agents']],[6.3,6.5,4.1],'角色及实际消息数量')
    largest=max(data['agents'],key=lambda a:a['message_count'])
    paragraph(f'结果中，{largest["label"]} 发出的消息最多，为 {largest["message_count"]} 条。Judge 的消息相对较少，但消息少不等于作用小，后面还需要结合具体记录看它是否参与合规判断。')

    page('五 通信渠道的分布')
    paragraph('读取 channel 后，可以把沟通分成内部讨论和公开发布两类。这里将 official_post、personal_post、anonymous_post 作为公开渠道，其余三个已知渠道作为内部沟通。')
    figure('03_channels.png','六种真实通信渠道的消息数量')
    channels=[]
    with (ROOT/'data/reports/channel_statistics.csv').open(encoding='utf-8-sig') as stream:channels=list(csv.DictReader(stream))
    table(['渠道','数量','占全部消息比例'],[[c['channel'],c['message_count'],f'{float(c["ratio"])*100:.2f}%'] for c in channels],[8.5,3.3,5.1],'渠道统计明细')
    paragraph(f'内部沟通合计 {total-public} 条，公开发布 {public} 条，约占全部消息的 {public/total*100:.2f}%。虽然公开消息数量不多，但它们直接面向外部，因此在后续筛选中需要单独观察。')

    page('六 轮次顺序与真实时间')
    paragraph('处理时间时发现，前期记录主要按天给出，危机当天改成按小时给出。如果直接把第 1 轮到第 23 轮当成连续的 23 小时，时间分布和通信速率都会出错。所以，我同时保留轮次顺序与原始时间戳。')
    figure('evidence/13_round_gaps.png','相邻轮次起点之间的实际间隔')
    figure('04_temporal.png','按真实时间定位的轮次消息分布')
    gaps=evidence['nonterminal_gap_counts']
    paragraph('相邻起点间隔统计为：'+ '，'.join(f'{float(k):g} 小时有 {v} 次' for k,v in gaps.items())+'。图中的空白仅说明没有给出新的轮次样本，不能解释成这一段时间完全没有通信。')
    paragraph('最后一轮没有下一轮起点，分析窗口按 18:00 到最后一条消息 18:55 计算，为 55 分钟。所有时间都保留官方未指定时区的模拟时间，不换算为当前真实日期。')
    paragraph('中间证据：reports/evidence/03_round_windows.csv 保存全部轮次起点、下轮起点、窗口时长、消息数和公开消息数。',style='Caption')

    page('七 缺失值与空值检查')
    paragraph('在做缺失值检查时，我把“字段不存在”“字段值是 null”和“数组为空”分开统计。这三种情况含义不同，不能直接合并后统一删除。')
    figure('evidence/14_field_quality.png','逐字段检查缺失 null 与空数组')
    table(['检查项','实际数量','本次处理'],[['缺少必要字段',sum(q['missing_fields'].values()),'无缺字段记录'],['internal_state 为 null',q['null_fields']['internal_state'],'保留消息，不补造内部状态'],['responding_to 为 null',q['null_fields']['responding_to'],'保留消息，不补造回复对象'],['recipients 为空数组',q['empty_recipients'],'保留为空，不猜测受众'],['不同 ID 的相同文本冗余',q['same_text_repetitions'],'保留各条 ID 和来源']],[6.2,2.1,8.6],'主要数据质量结果及处理方式')
    paragraph('internal_state 为空，只能说明这条消息没有给出内部状态；它本身的正文、发送者和时间仍然完整。如果将这类消息删除，会丢失大量正常通信记录。因此，本次质量检查以标记和保留为主。')

    page('八 数据过滤与保留策略')
    paragraph('数据过滤主要用于缩小分析范围，并没有把所有与并购无关的消息从底表中删掉。正常消息还要作为行为对照，所以处理后的消息表保留全部原始消息。')
    figure('07_filtering.png','原始消息 标准化保留与分析子集的对比')
    table(['处理情况','采用的规则'],[['完全重复 ID','同 ID 同内容才去重；同 ID 内容冲突则报错'],['文本重复但 ID 不同','不删除，保留每次记录'],['关键词筛选','只产生主题分析子集，不直接判为泄露'],['公开渠道筛选','按真实 channel 分类，不猜测发布对象'],['无效和未来回复','保留原值与质量标记，不进入默认传播线索']],[5.4,11.5],'本次过滤和清洗规则')
    paragraph(f'原始与标准化消息均为 {total} 条，没有因清洗排除消息。主题关键词子集为 {sum(m["topic_match"] for m in data["messages"])} 条，公开渠道子集为 {public} 条，异常候选支持消息为 {evidence["candidate_support_messages"]} 条。后三个子集互有重叠，不应把它们看成逐层减少的流水线。')

    page('九 消息抽取前后的对照')
    paragraph('为了检查字段整理有没有改变原始含义，我保留了几条完整的前后对照。下面以一条 side_huddle 消息为例，只展示与结构转换有关的字段，完整正文仍保留在中间 JSON 文件中。')
    sample=examples[1];raw=sample['raw_record'];normalized=sample['normalized_record']
    code(json.dumps({key:raw[key] for key in ['message_id','agent_id','timestamp','channel','recipients','responding_to']},ensure_ascii=False,indent=2))
    table(['处理后字段','本条实际值'],[['message_id',normalized['message_id']],['round_index',normalized['round_index']],['scope',normalized['scope']],['recipients','\n'.join(normalized['recipients'])],['source_file',normalized['source_file']],['json_pointer',normalized['json_pointer']]],[4.2,12.7],'同一条消息的标准化结果')
    paragraph('这里新增了 round_index、scope 和来源定位字段，原始 message_id 与正文没有改写。通过 JSON Pointer，可以回到第 14 轮的第 4 条原始消息检查。程序内部从 0 开始编号，所以表中的 round_index 为 13。')
    paragraph('中间证据：reports/evidence/06_before_after_examples.json 同时保存四条原始记录、对应标准化记录、生成的边和内部状态。',style='Caption')

    page('十 接收人名称规范化与关系展开')
    paragraph('另一个需要处理的问题是接收人名称。发送者使用 agent_id，但 recipients 中经常出现 legal、pr、platform_trust 等角色名。这里根据原数据中的 ID、role 和 label 做无歧义映射，再生成通信关系。')
    table(['原始接收人','规范化结果'],[[row['raw_recipient'],row['normalized_agent_id'] or '保留广播标志，不展开节点'] for row in evidence['role_mapping']],[5.2,11.7],'接收人映射表')
    figure('evidence/23_edge_example.png','消息 20460605_13_004 生成的三条指定接收人边')
    paragraph('本条消息原始接收人是 platform_trust、pr 和 social_manager，处理后分别映射到 quality_agent、pr_agent 和 social_media_agent。因此，一条消息可以支持三条接收人关系。')
    paragraph('ALL 只保留为广播标志，不生成名叫 ALL 的虚构节点，也不把所有 Agent 都写成已经阅读了消息。这个区别会直接影响后面的网络图。')

    page('十一 回复引用的完整性检查')
    paragraph('最开始查看 responding_to 时，看起来它都像是“回复给谁”的字段，但逐条检查后发现，其中既有消息 ID，也有 @角色文本，还有指向未来记录的 ID。这里需要分别处理，不能只要字段非空就建立同一种边。')
    figure('evidence/15_reply_quality.png','全部消息的回复字段分类结果')
    future=examples[2]['normalized_record'];parent=next(m for m in data['messages'] if m['message_id']==future['responding_to'])
    table(['实例','原始记录','检查结果'],[['@文本引用',sample['message_id']+'\n'+raw['responding_to'],'不能解析为消息 ID，保留审计'],['未来消息引用',future['message_id']+'\n'+future['timestamp']+' → '+parent['timestamp'],'目标记录存在，但时间晚于当前消息']],[3.1,9.8,4.0],'两种问题引用的真实实例')
    paragraph(f'有效回复为 {evidence["reply_counts"]["时间有效引用"]} 条，未来引用 {evidence["reply_counts"]["未来引用"]} 条，非消息 ID 引用 {evidence["reply_counts"]["非消息ID引用"]} 条，空引用 {evidence["reply_counts"]["空引用"]} 条。未来引用保留 temporal_valid=false，不参与默认网络和传播线索。')
    paragraph(f'中间证据：reports/evidence/05_reply_reference_audit.csv 为全部 {total} 条消息逐条保留了引用值、目标时间、分类和原始位置。',style='Caption')

    page('十二 标准化数据表与保存结果')
    paragraph('完成字段整理后，我没有只留下一个前端 JSON，而是同时保存八张 CSV。这样后续做不同图表或改后端时，都可以直接读取相应的数据表。')
    figure('evidence/16_tables.png','八张输出表的记录数量')
    units={'messages':'一条消息','agents':'一个角色','events':'一条环境条目','internal_states':'一条非空状态流','communication_edges':'一条接收或回复边','participants':'角色与轮次参与记录','rounds':'一个轮次','candidates':'一次规则候选'}
    table(['表名','行数','统计单元'],[[r['table'],r['rows'],units[r['table']]] for r in evidence['table_inventory']],[7.0,2.2,7.7],'结构化输出清单')
    paragraph('表之间的行数不同是正常的。例如一条消息可以展开成多条通信边，也可以包含多条内部状态。不能把这些表的行数简单相加后当作消息总量。每张表前三条记录还单独保存为 preview_*.json，方便快速查看输出格式。')

    page('十三 内部状态与外部环境')
    paragraph('除了消息正文，内部状态和环境信息也需要保留。内部状态是 Agent 记录的反应、解释或思考，环境条目则包含新闻、外部行动和市场背景。两者的来源层级不同，因此分别建表。')
    figure('evidence/17_internal_states.png','非空内部状态的三种类型')
    figure('evidence/18_events.png','环境上下文中实际提取的事件类型')
    paragraph(f'内部状态共 {len(data["internal_states"])} 条，其中 deliberating 为 {evidence["internal_state_types"]["deliberating"]} 条。环境事件共 {len(data["events"])} 条。环境条目使用轮次起点作为时间，并标注轮次精度；正文中提到的预测时刻不另行伪造为精确事件。')

    page('十四 Agent 行为画像与速率')
    paragraph('得到消息表后，可以进一步比较各 Agent 的沟通方式。这里除了统计总消息量，也分别查看内部与公开渠道比例，避免只根据“谁说话最多”来判断风险。')
    figure('evidence/19_agent_scope.png','不同 Agent 的内部消息与公开发布')
    figure('evidence/21_volume_rate.png','同一批轮次的消息量和按窗口时长计算的速率')
    paragraph('下图说明了为什么需要保留真实时间。前期一轮可能覆盖 24 或 72 小时，危机当天一轮主要覆盖 1 小时。按窗口除出的速率是样本内指标，并不能代表这段时间所有实际通信都被连续记录。总体速率统一使用全语料窗口总时长作为分母，以便比较。')

    page('十五 异常候选规则与初步结果')
    paragraph('本次先使用可以解释的规则寻找候选，不把模型分数直接当成结论。规则考虑通信速率、新的公开渠道、实体公开，以及人员不可用时发生的主题决策。每次命中都保留对应消息及触发原因。')
    figure('evidence/20_rules.png','五类规则的候选数量')
    rule_labels={'rate_surge':'轮次速率跃升','decision_during_unavailability':'人员不可用时的主题决策','named_public_topic':'公开实体主题','named_public_before_formal_lift':'正式解禁前实体公开','new_public_channel':'首次公开渠道'}
    table(['规则','候选数','支持消息数'],[[rule_labels[r['rule']],r['candidates'],r['support_messages']] for r in evidence['rule_stats']],[10.0,3.0,3.9],'候选与支持消息的区别')
    paragraph(f'最终得到 {len(data["candidates"])} 条候选，关联 {evidence["candidate_support_messages"]} 条不同的支持消息。这里候选数和消息数不能混用：整轮速率跃升对应一轮内的一组消息，而不是只有作为定位点的第一条消息。')
    paragraph('例如，公开实体关键词命中后，还需要检查是否在正式解禁前、是否只是转述、以及是否存在授权记录。本次的风险分数仅用于安排人工查看顺序，不是经过验证的泄露概率。')

    page('十六 HarborCrest 关键记录核对')
    paragraph('为了把统计结果与具体事件联系起来，我从数据中保留了几条有代表性的原始记录。下面按记录时间排列，完整英文原文、来源文件和 JSON Pointer 均保存在中间证据文件中。')
    figure('evidence/22_key_records.png','从内部提及到公开与正式解禁的记录顺序')
    table(['记录时间','消息 ID','这条记录能说明什么'],[[r['timestamp'][5:].replace('T',' '),r['message_id'],r['kind']] for r in markers],[3.6,6.0,7.3],'关键记录的可追溯定位')
    first=data['evidence_summary']['first_named_public'];release=data['evidence_summary']['formal_lift']
    paragraph(f'首次直接包含交易实体的公开消息是 {first["message_id"]}，发送时间为 {first["timestamp"][11:16]}；明确写出正式解禁的记录是 {release["message_id"]}，时间为 {release["timestamp"][11:16]}。这两个时间来自原消息，不是在程序中预设的。')
    paragraph('同时，17:19 的 Legal 消息声称取得了口头双边同意并下达 GO 指令。这属于数据中记录的授权声称，需要与正式解禁标记一起阅读，不能因为关键词或时间差就直接认定违规。最早的主题线索还包含推测性内部状态，也不能当作已经获知真实交易信息。')
    paragraph('中间证据：reports/evidence/11_key_evidence_records.json；reports/harborcrest_findings.md。',style='Caption')

    page('十七 可视化映射结果')
    paragraph('经过前面的处理，热力图可以使用 Agent、轮次和消息数，网络图可以使用标准化通信边。下面两张图都直接读取处理结果，不手工添加节点或调整统计值。')
    figure('05_heatmap.png','Agent 与轮次对应的消息数量',width=6.35)
    figure('06_network.png','基于有效时间记录构建的 Agent 通信网络',width=5.6)
    paragraph(f'网络底表共有 {len(data["communication_edges"])} 条关系，其中时间有效 {evidence["valid_edge_count"]} 条；再隐藏 {evidence["hidden_self_edges"]} 条自环后，图中使用 {evidence["network_edge_records"]} 条关系记录，合并成 {evidence["network_aggregate_edges"]} 组方向关系。箭头表示收件或回复联系，不单独证明敏感信息已经传播。')

    page('十八 运行验证与实际页面')
    paragraph('完成处理后，我又检查了一遍输出：CSV 行数是否与内存统计一致、图表总数是否能回算，以及消息是否能通过来源指针找到原始正文。这样可以避免程序虽然运行成功，但生成结果之间对不上的情况。')
    table(['本次重新执行的命令','退出码'],[[r['command'],r['exit_code'] if r['exit_code'] is not None else '未执行'] for r in verification],[14.7,2.2],'实际验证执行记录')
    paragraph(f'测试使用标注的结构样本，并在本地官方数据存在时检查真实统计一致性；当前 {passed_tests} 项 Python 与 API 检查通过，来源验证逐条核对 {total} 条消息。验证日志原样保存在中间证据目录。')
    screenshot=ROOT/'reports/figures/08_dashboard_overview.png'
    if screenshot.exists():figure('08_dashboard_overview.png','四图版前端页面实际运行画面 局部',width=4.2)
    paragraph('前端目前只保留行为热力图、通信网络、事件时间轴和渠道分布图，每张图配简单说明。数据处理细节放在本报告和中间文件中，页面主要用于观察图表。')

    page('十九 实验中遇到的问题与小结')
    heading('数据结构与文档示例不一致')
    paragraph('最开始如果直接按照说明中的 agent_outputs 编写解析器，就不能正确读出实际消息。后来先检查真实 JSON，再按 rounds 下的 communications 抽取，才把记录完整读出来。这个过程说明，数据说明可以帮助理解背景，但不能代替对实际文件的检查。')
    heading('时间和回复字段不能只看表面格式')
    paragraph('这次比较容易忽略的地方是轮次间隔和回复引用。轮次看起来是顺序编号，实际却混合了按天和按小时的记录；responding_to 看起来是引用字段，实际还混有角色文本和未来 ID。把这两类情况保存成中间表以后，问题就比较容易复查。')
    heading('候选与结论需要区分')
    paragraph('从图表中可以较快发现公开行为和通信量变化，但要解释为什么发生，还需要继续读原文。本次先把规则依据保存下来，没有直接把候选定性为真实泄露。后续需要结合授权、预期角色行为和人工标注来完善。')
    heading('本次实验小结')
    paragraph('通过本次实验，数据下载、结构理解、过滤和标准化处理已经形成了一条可重复运行的流程。相比只保留最终图表，现在还保留了逐字段检查、完整回复审计、名称映射和前后对照实例，后面修改分析方法时也有依据。')
    paragraph('目前做得还不够的是：关键词筛选可能漏掉间接表达，速率规则容易受到抽样时间窗口影响，授权信息的真伪也不能只依靠消息中的自述。这些内容会作为后续改进重点。学生姓名、学号和课程信息仍需补充。')
    paragraph('可复现的主要命令如下：')
    code('python scripts/download_data.py\npython scripts/preprocess.py\npython scripts/build_evidence.py\npython scripts/generate_evidence_figures.py\npython scripts/record_evidence_checks.py\npython scripts/write_evidence_report.py\npython scripts/verify_evidence.py')

    page('附录 中间证据与复查位置')
    paragraph('下表列出报告中使用的主要中间证据。原始大型 JSON 不入库，保存在 data/raw；统计、审计表和少量可定位的实例保留在 reports/evidence。')
    files=[('01_archive_inventory.csv','压缩包内全部文件和大小'),('02_message_field_audit.csv','消息字段完整性与空值统计'),('03_round_windows.csv','全部轮次时间及窗口时长'),('04_recipient_mapping.csv','接收人名称规范化映射'),('05_reply_reference_audit.csv','全部消息的回复引用检查'),('06_before_after_examples.json','四条真实记录的前后对照'),('07_normalized_table_inventory.csv','八张标准表的行数和列数'),('08_filter_subsets.csv','原始与分析子集的计数'),('09_rule_statistics.csv','各规则候选及支持集合规模'),('10_agent_behavior.csv','角色内部与公开行为画像'),('11_key_evidence_records.json','五条关键记录的完整原文'),('12_verification_runs.json 和 txt','实际验证命令与原样输出'),('preview_*.json','每张标准表的前三条样本'),('summary.json','本报告使用的汇总值')]
    table(['reports/evidence 下的文件','保留的内容'],files,[9.7,7.2],'中间证据文件索引')
    paragraph('新增图表在 reports/figures/evidence 下，manifest.json 保存每张图引用的统计值。原有七张课程图表在 reports/figures 下。对照消息的完整正文与 ID 可以回到官方 JSON 的对应位置，不需要依赖报告中的节选。')
    paragraph('数据和任务来源为 IEEE VAST Challenge 2026 MC1。代码与统计由当前项目重新生成。公开记录时间与记录中声称的授权需要分别核验，本文不将它们合并成未经证明的因果结论。',style='Caption')

    for p in doc.paragraphs:
        for run in p.runs:
            if run.font.name != 'Courier New':run.font.name='Arial Unicode MS'
            font=run._element.get_or_add_rPr().get_or_add_rFonts()
            for attr in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:font.attrib.pop(qn('w:'+attr),None)
            font.set(qn('w:eastAsia'),'Arial Unicode MS')
    doc.core_properties.title='多 Agent 异常行为与信息传播可视分析 数据处理实验报告'
    doc.core_properties.author='课程小组 待填写'
    OUTPUT.parent.mkdir(parents=True,exist_ok=True);doc.save(OUTPUT)
    manifest={'file':str(OUTPUT.relative_to(ROOT)),'figures':figure_count,'tables':table_count,'outline':outline,'figure_sources':source_index,
              'reference_voice':'按实验目的 步骤 运行结果 结果分析 小结展开；采用自然的课程实验叙述，不复制参考身份经历或实验结论',
              'data':{'messages':total,'agents':len(data['agents']),'rounds':len(data['rounds']),'candidates':len(data['candidates'])}}
    (EVIDENCE/'report_manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2))
    print(json.dumps({'output':str(OUTPUT),'figures':figure_count,'tables':table_count,'sections':len(outline)+1},ensure_ascii=False))

if __name__=='__main__':create()
