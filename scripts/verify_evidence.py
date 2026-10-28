"""Cross-check the retained evidence pack against original records and report sources."""
import csv
import json
from collections import Counter
from pathlib import Path
from docx import Document
from inspect_data import ROOT


def verify():
    evidence=ROOT/'reports/evidence';data=json.loads((ROOT/'data/processed/dashboard.json').read_text())
    summary=json.loads((evidence/'summary.json').read_text());raw=json.loads((ROOT/summary['source_file']).read_text())
    messages=[message for round_ in raw['rounds'] for message in round_['communications']]
    assert summary['message_count']==len(messages)
    assert sum(summary['reply_counts'].values())==len(messages)
    raw_by_id={message['message_id']:message for message in messages}
    expected_refs=Counter()
    for message in messages:
        ref=message.get('responding_to');parent=raw_by_id.get(ref)
        category='空引用' if ref is None else '非消息ID引用' if parent is None else '未来引用' if parent['timestamp']>message['timestamp'] else '时间有效引用'
        expected_refs[category]+=1
    assert dict(expected_refs)==summary['reply_counts']
    for row in summary['fields']:
        assert row['present']+row['missing']==len(messages)
        assert row['null']==sum(row['field'] in message and message[row['field']] is None for message in messages)
    examples=json.loads((evidence/'06_before_after_examples.json').read_text())
    for example in examples:
        original=raw
        for key in example['json_pointer'].strip('/').split('/'):
            original=original[int(key)] if isinstance(original,list) else original[key]
        assert original==example['raw_record']
        assert example['normalized_record']==next(message for message in data['messages'] if message['message_id']==example['message_id'])
    with (evidence/'03_round_windows.csv').open(encoding='utf-8-sig') as stream:windows=list(csv.DictReader(stream))
    assert len(windows)==len(raw['rounds'])
    assert sum(int(row['messages']) for row in windows)==len(messages)
    for row in summary['table_inventory']:
        with (ROOT/row['path']).open(encoding='utf-8-sig') as stream:count=len(list(csv.DictReader(stream)))
        assert count==row['rows']==len(data[row['table']])
    assert summary['candidate_support_messages']==len({mid for c in data['candidates'] for mid in c['supporting_message_ids']})
    figures=json.loads((ROOT/'reports/figures/evidence/manifest.json').read_text())['figures']
    assert len(figures)==14
    for figure in figures:assert (ROOT/figure['file']).stat().st_size>1000
    by_name={Path(figure['file']).stem:figure for figure in figures}
    assert by_name['15_reply_quality']['totals']==summary['reply_counts']
    assert by_name['16_tables']['totals']=={row['table']:row['rows'] for row in summary['table_inventory']}
    assert by_name['17_internal_states']['totals']==dict(Counter(state['state_type'] for state in data['internal_states']))
    assert by_name['18_events']['totals']==dict(Counter(event['type'] for event in data['events']))
    assert sum(by_name['20_rules']['totals'].values())==len(data['candidates'])
    assert sum(by_name['21_volume_rate']['totals']['volumes'])==len(messages)
    report=json.loads((evidence/'report_manifest.json').read_text());document=Document(ROOT/report['file'])
    assert len(document.inline_shapes)==report['figures']==21
    assert len(document.tables)==report['tables']==16
    for figure in report['figure_sources']:assert (ROOT/figure['figure']).exists()
    paragraphs='\n'.join(paragraph.text for paragraph in document.paragraphs)
    assert '8209190423' not in paragraphs and '卓世龙' not in paragraphs
    assert not any(border for style in document.styles for border in style.element.iter('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}pBdr'))
    result={'status':'passed','raw_messages':len(messages),'source_examples':len(examples),'round_windows':len(windows),
            'tables':len(summary['table_inventory']),'additional_figures':len(figures),'word_figures':len(document.inline_shapes),'word_tables':len(document.tables)}
    (evidence/'verification.json').write_text(json.dumps(result,ensure_ascii=False,indent=2));print(json.dumps(result,ensure_ascii=False))

if __name__=='__main__':verify()
